from fastapi import FastAPI, Body, File, UploadFile
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.requests import Request
from fastapi.responses import FileResponse
import psycopg2
import psycopg2.extras
import pandas as pd
import os
import tempfile
from datetime import date,datetime
from dotenv import load_dotenv
import tempfile
from starlette.background import BackgroundTask
import openpyxl
from openpyxl.utils import get_column_letter

load_dotenv()

STATUS_COLUMNS = [
    "Development",
    "KLD",
    "Artwork",
    "Trial",
    "Implementation",
    "Comments",
    "Status"
]

app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

def get_conn():
    return psycopg2.connect(
        host=os.getenv("DB_HOST"),
        port=os.getenv("DB_PORT"),
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD")
    )

def dict_cursor(conn):
    # psycopg2 equivalent of sqlite3's row_factory
    return conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

@app.get("/")
def index(request: Request):
    return templates.TemplateResponse(request, "index.html")

@app.get("/value/api/projects")
def get_projects():
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute("SELECT project_name,MIN(id) FROM projects GROUP BY project_name ORDER BY MIN(id) ASC")
    rows = cur.fetchall()
    cur.close()
    conn.close()
    # print(f"Projects : {[row["project_name"] for row in rows]}")
    return [row["project_name"] for row in rows]

@app.get("/value/api/projects/id/{project_id}")
def get_project_by_id(project_id: int):
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute("SELECT * FROM projects WHERE id = %s", (project_id,))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return dict(row)

@app.get("/value/api/projects/{project_name}")
def get_project_rows(project_name: str):
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute("""
        SELECT p.id, p.pm_code, p.project_name, p.vendor, p.packaging_type, p.packaging_option, p.eta
        FROM projects p WHERE p.project_name = %s ORDER BY p.id ASC
    """, (project_name,))
    rows = cur.fetchall()

    result = []
    today = date.today().isoformat()
    TEXT_COLS = ["Comments"]

    for row in rows:
        row_dict = dict(row)
        project_id = row_dict["id"]

        cur.execute("SELECT column_name, current_value FROM status WHERE project_id = %s", (project_id,))
        statuses = cur.fetchall()

        cur.execute("SELECT column_name, deadline FROM deadlines WHERE project_id = %s", (project_id,))
        deadlines = cur.fetchall()

        status_map = {s["column_name"].strip(): s["current_value"] for s in statuses}
        deadline_map = {d["column_name"].strip(): (str(d["deadline"]) if d["deadline"] is not None else None) 
                        for d in deadlines}
        # str() needed — psycopg2 returns deadline as a Python date object, not string

        for col in STATUS_COLUMNS:
            row_dict[col] = status_map.get(col, None)
        TEXT_COLS=["Comments"]

        red_cols = []
        yellow_cols = []
        for col, deadline in deadline_map.items():
            if col in TEXT_COLS:
                continue
            current_val = status_map.get(col, "")
            is_complete = current_val in ["Completed","Shared"]
            is_overdue = deadline and today > deadline

            if is_overdue and not is_complete:
                red_cols.append(col)
            elif is_overdue and is_complete:
                yellow_cols.append(col)

        row_dict["red_cols"] = red_cols
        row_dict["yellow_cols"] = yellow_cols
        row_dict["_health"] = "green"
        result.append(row_dict)

    cur.close()
    conn.close()
    return result

@app.get("/value/api/status/{project_id}")
def get_status(project_id: int):
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute("SELECT * FROM status WHERE project_id = %s", (project_id,))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [dict(row) for row in rows]

@app.get("/value/api/deadlines/{project_id}")
def get_deadlines(project_id: int):
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute("SELECT column_name, deadline FROM deadlines WHERE project_id = %s", (project_id,))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    # Convert date objects to strings for JSON serialisation
    return [{"column_name": r["column_name"], "deadline": str(r["deadline"]) if r["deadline"] else ""} for r in rows]

@app.get("/value/api/eta/{project_id}")
def get_eta(project_id:int):
    conn=get_conn()
    cur=dict_cursor(conn)
    cur.execute("SELECT eta FROM projects where id=%s",(project_id,))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return {"project_id": project_id, "eta": str(row["eta"]) if row["eta"] else ""}

@app.get("/value/api/alerts")
def get_alerts():
    conn = get_conn()
    cur = dict_cursor(conn)
    today = date.today().isoformat()
    cur.execute("""
        SELECT p.id as project_id, p.project_name, p.packaging_type, p.packaging_option,p.vendor,p.eta,
               s.column_name, s.current_value, d.deadline
        FROM deadlines d
        JOIN projects p ON p.id = d.project_id
        JOIN status s ON s.project_id = d.project_id AND s.column_name = d.column_name
        WHERE d.deadline < %s
        AND (s.current_value IS NULL OR s.current_value NOT IN ('Completed','Shared'))
        ORDER BY d.deadline ASC
    """, (today,))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [dict(r) | {"deadline": str(r["deadline"])} for r in rows]

@app.get("/value/api/alerts/{project_name}")
def get_alerts_for_project(project_name: str):
    conn = get_conn()
    cur = dict_cursor(conn)
    today = date.today().isoformat()
    cur.execute("""
        SELECT p.id as project_id, p.project_name, p.packaging_type, p.packaging_option, p.vendor, p.eta,
               s.column_name, s.current_value, d.deadline
        FROM deadlines d
        JOIN projects p ON p.id = d.project_id
        JOIN status s ON s.project_id = d.project_id AND s.column_name = d.column_name
        WHERE d.deadline < %s AND p.project_name = %s
        AND (s.current_value IS NULL OR s.current_value NOT IN ('Completed','Shared'))
        ORDER BY d.deadline ASC
    """, (today, project_name))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [dict(r) | {"deadline": str(r["deadline"])} for r in rows]



@app.get("/value/api/due-today")
def get_due_today():
    conn = get_conn()
    cur = dict_cursor(conn)
    today = date.today().isoformat()
    cur.execute("""
        SELECT p.id as project_id, p.project_name, p.packaging_type, p.packaging_option, p.vendor, p.eta,
               s.column_name, s.current_value, d.deadline
        FROM deadlines d
        JOIN projects p ON p.id = d.project_id
        JOIN status s ON s.project_id = d.project_id AND s.column_name = d.column_name
        WHERE d.deadline = %s
        AND (s.current_value IS NULL OR s.current_value NOT IN ('Completed','Shared'))
        ORDER BY p.project_name ASC
    """, (today,))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [dict(r) | {"deadline": str(r["deadline"])} for r in rows]

@app.get("/value/api/due-today/{project_name}")
def get_due_today_for_project(project_name: str):
    conn = get_conn()
    cur = dict_cursor(conn)
    today = date.today().isoformat()
    cur.execute("""
        SELECT p.id as project_id, p.project_name, p.packaging_type, p.packaging_option, p.vendor, p.eta,
               s.column_name, s.current_value, d.deadline
        FROM deadlines d
        JOIN projects p ON p.id = d.project_id
        JOIN status s ON s.project_id = d.project_id AND s.column_name = d.column_name
        WHERE d.deadline = %s AND p.project_name = %s
        AND (s.current_value IS NULL OR s.current_value NOT IN ('Completed','Shared'))
        ORDER BY p.project_name ASC
    """, (today, project_name))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [dict(r) | {"deadline": str(r["deadline"])} for r in rows]

@app.put("/value/api/status/{project_id}")
def update_status(project_id: int, data: dict = Body(...)):
    conn = get_conn()
    cur = conn.cursor()
    today = date.today().isoformat()
    GREEN_VALUES = ["Completed","Shared"]

    for col, value in data.items():
        completion_date = today if value in GREEN_VALUES else None
        cur.execute("""
            INSERT INTO status (project_id, column_name, current_value, completion_date)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT(project_id, column_name) DO UPDATE SET
                current_value = EXCLUDED.current_value,
                completion_date = CASE
                    WHEN EXCLUDED.current_value IN ('Completed','Shared')
                    THEN COALESCE(status.completion_date, EXCLUDED.completion_date)
                    ELSE NULL
                END
        """, (project_id, col, value, completion_date))
    # Note: PostgreSQL ON CONFLICT syntax uses EXCLUDED.column instead of bare values

    conn.commit()
    cur.close()
    conn.close()
    return {"status": "ok"}

@app.put("/value/api/projects/{project_id}/eta")
def update_eta(project_id: int, data: dict = Body(...)):
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute(
        "UPDATE projects SET eta=%s WHERE id=%s",
        (data["eta"] or None, project_id)
    )
    conn.commit()
    cur.close()
    conn.close()
    return {"status": "ok"}


@app.post("/value/api/deadlines/{project_id}")
def save_deadlines(project_id: int, data: dict = Body(...)):
    conn = get_conn()
    cur = conn.cursor()
    for col, deadline in data.items():
        if deadline:
            cur.execute("""
                INSERT INTO deadlines (project_id, column_name, deadline)
                VALUES (%s, %s, %s)
                ON CONFLICT(project_id, column_name) DO UPDATE SET deadline = EXCLUDED.deadline
            """, (project_id, col, deadline))
    conn.commit()
    cur.close()
    conn.close()
    return {"status": "ok"}

@app.post("/value/api/projects")
def add_project(data: dict = Body(...)):
    conn = get_conn()
    cur = conn.cursor()
    eta_val = datetime.strptime(data["ETA"], "%Y-%m-%d").date() if data.get("ETA") else None
    pm_code = data.get("PM Code")
    project_desc = data.get("Project")
    composite_name = f"PM - {pm_code} - {project_desc}" if pm_code and project_desc else project_desc
    cur.execute("""
        INSERT INTO projects (project_name, packaging_type, packaging_option,vendor,pm_code,eta)
        VALUES (%s, %s, %s, %s, %s, %s) RETURNING id
    """, (composite_name, data["Packaging Type"], data["Packaging Option"],data["Vendor"],data["PM Code"],eta_val))
    project= cur.fetchone()
    # print(project)
    project_id=project[0]
    for col in STATUS_COLUMNS:
        cur.execute("""
            INSERT INTO status (project_id, column_name, current_value)
            VALUES (%s, %s, NULL) RETURNING project_id
        """, (project_id, col))
    conn.commit()
    cur.close()
    conn.close()
    return {"status": "ok"}

@app.delete("/value/api/projects/{project_id}")
def delete_project(project_id: int):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("DELETE FROM projects WHERE id = %s", (project_id,))
    conn.commit()
    cur.close()
    conn.close()
    return {"status": "ok"}


@app.get("/value/api/download/{excel_name}")
def download_excel(excel_name: str):
    TEXT_COLS=["Comments"]
    conn = get_conn()
    cur = dict_cursor(conn)
    cur.execute("SELECT DISTINCT project_name,MIN(id) FROM projects GROUP BY project_name ORDER BY MIN(id) ASC")
    projects = [r["project_name"] for r in cur.fetchall()]
    all_rows = []

    for project in projects:
        cur.execute("""
            SELECT p.id, p.pm_code, p.vendor, p.project_name, p.packaging_type, p.packaging_option, p.eta
            FROM projects p WHERE p.project_name = %s
        """, (project,))
        rows = cur.fetchall()
        # print(rows)
        for row in rows:
            project_id = row["id"]
            cur.execute("SELECT column_name, current_value, completion_date FROM status WHERE project_id = %s", (project_id,))
            statuses = cur.fetchall()
            cur.execute("SELECT column_name, deadline FROM deadlines WHERE project_id = %s", (project_id,))
            deadlines = cur.fetchall()

            status_map = {s["column_name"]: s["current_value"] for s in statuses}
            completion_map = {s["column_name"]: str(s["completion_date"]) if s["completion_date"] else None for s in statuses}
            deadline_map = {d["column_name"]: str(d["deadline"]) if d["deadline"] else None for d in deadlines}

            flat = {
                "S. No.":row["id"],
                "Project": row["project_name"],
                "PM Code":row["pm_code"],
                "Packaging Type": row["packaging_type"],
                "Packaging Option": row["packaging_option"],
                "Vendor":row["vendor"],
                "ETA":row["eta"]
            }
            for col in STATUS_COLUMNS:
                flat[col] = status_map.get(col)
                if(col not in TEXT_COLS):
                    flat[col + " | Deadline"] = deadline_map.get(col)
                    flat[col + " | Completed On"] = completion_map.get(col)
            all_rows.append(flat)

    cur.close()
    conn.close()

    df = pd.DataFrame(all_rows)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp_path = tmp.name

    df.to_excel(tmp_path, index=False, sheet_name="VE Tracker")

    # Auto-fit column widths
    wb = openpyxl.load_workbook(tmp_path)
    ws = wb.active
    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col_cells
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 4, 60)
        # +4 for padding, capped at 60 so very long text columns don't become huge
    wb.save(tmp_path)
        
    return FileResponse(
        path=tmp_path,
        filename=f"{excel_name}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        background=BackgroundTask(os.remove, tmp_path)
    )

@app.post("/value/api/import")
async def import_excel(file: UploadFile = File(...)):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        contents = await file.read()
        tmp.write(contents)
        tmp_path = tmp.name

    xl = pd.ExcelFile(tmp_path)
    sheet = xl.sheet_names[0]
    df = xl.parse(sheet_name=sheet)#header = 1 for the original excel but for the rest its header=0, so we can just use pandas default of header=0 which treats the first row as header. If the uploaded excels have a different format, we can add logic to detect and handle that.
    xl.close()
    # print(f"Columns : {df.columns}")
    if "Project" in df.columns:
        df["Project"] = df["Project"].ffill()

    if "PM Code" in df.columns:
        df["PM Code"] = df["PM Code"].apply(
            lambda x: str(int(x)) if pd.notna(x) else None
        )

    if "ETA" in df.columns:
        df["ETA"] = pd.to_datetime(df["ETA"], errors='coerce')
    # print(f"DF currently is {df.head()}")
    conn = get_conn()
    cur = conn.cursor()
    rows_imported = 0
    GREEN_VALUES = ["Completed","Shared"]

    for _, row in df.iterrows():
        raw_eta = row.get("ETA")
        eta_val = raw_eta.to_pydatetime().date() if pd.notna(raw_eta) and hasattr(raw_eta, 'to_pydatetime') else None

        pm_code = row.get("PM Code")
        project_desc = row.get("Project")
        if pd.isna(project_desc) or str(project_desc).strip() == "":
            continue
        # composite_name = f"PM - {pm_code} - {project_desc}" if pm_code and project_desc else project_desc
        # print(f"Simple Name:{project_desc}")
        cur.execute("""
            INSERT INTO projects (project_name, packaging_type, packaging_option, vendor, pm_code, eta)
            VALUES (%s, %s, %s, %s, %s, %s) RETURNING id
        """, (
            project_desc, 
            row.get("Packaging Type"), 
            row.get("Packaging Option"), 
            row.get("Vendor"), 
            pm_code, 
            eta_val
        ))
        project_id = cur.fetchone()[0]

        for col in STATUS_COLUMNS:
            value = row.get(col)
            value = str(value).title().strip() if pd.notna(value) else None

            completion_date = None
            completion_col = col + " | Completed On"
            
            if completion_col in df.columns:
                raw_comp = row.get(completion_col)
                if pd.notna(raw_comp):
                    try:
                        completion_date = pd.to_datetime(raw_comp).date().isoformat()
                    except Exception:
                        pass

            # if completion_date is None and value in GREEN_VALUES:
            #     completion_date = date.today().isoformat()

            cur.execute("""
                INSERT INTO status (project_id, column_name, current_value, completion_date)
                VALUES (%s, %s, %s, %s)
            """, (project_id, col, value, completion_date))

            deadline_col = col + " | Deadline"
            if deadline_col in df.columns:
                raw_deadline = row.get(deadline_col)
                if pd.notna(raw_deadline):
                    try:
                        deadline_str = pd.to_datetime(raw_deadline).date().isoformat()
                        cur.execute("""
                            INSERT INTO deadlines (project_id, column_name, deadline)
                            VALUES (%s, %s, %s)
                            ON CONFLICT(project_id, column_name) DO UPDATE SET deadline = EXCLUDED.deadline
                        """, (project_id, col, deadline_str))
                    except Exception:
                        pass

        rows_imported += 1

    conn.commit()
    cur.close()
    conn.close()

    try:
        os.remove(tmp_path)
    except Exception:
        pass

    return {"status": "ok", "rows_imported": rows_imported}